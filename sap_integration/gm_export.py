"""Parser for SAP Grants Management line-item exports (GM_E_4GBA).

The export lists every booking of one PSP-Element over the whole project
runtime. Detail rows carry a Fonds; subtotal rows after each document and the
grand total row at the end do not and are skipped. Payments reference their
reservation/order through "Vorg. Referenzschl.", which lets us group all
bookings of one staff contract, trip or purchase into a single position.
"""

import re
from calendar import monthrange
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from django.db import transaction
from openpyxl import load_workbook

from projects.models import SAPFund
from sap_integration.models import SAPImport, SAPPosition, SAPPositionKind


REQUIRED_HEADERS = {
    "PSP-Element",
    "Fonds",
    "E/A-Art",
    "Geschäftsjahr",
    "Buchungsperiode",
    "Referenzbelegnummer",
    "Vorg. Referenzschl.",
    "Text",
    "Buchungsdatum",
    "Transaktionswährung",
    "FMM-Werttyp",
    "Belegkopftext",
}

BUDGET_VALUE_TYPES = {"R1"}
COMMITMENT_VALUE_TYPES = {"50", "51", "81"}
ACTUAL_VALUE_TYPES = {"66", "99", "Z1"}

# E/A-Art codes differ between projects (internal codes such as 7221 or the
# funder's positions such as 0812), so the kind of a position is primarily
# derived from the Finanzposition: in the budget classification its group
# 4xx is personnel, 2xx income. The codes below are only a fallback.
STAFF_BUDGET_GROUP = "4"
INCOME_BUDGET_GROUP = "2"
STAFF_COST_TYPES = {"7221", "PERSONALKOSTEN", "0812", "0817", "0822"}
TRAVEL_COST_TYPES = {"7464", "0846"}
INCOME_COST_TYPES = {"7999"}
TRAVEL_TEXT_RE = re.compile(r"^\W*RK[EA]\b|^\W*RK[EA]_")

MONTHS = {
    "jan": 1, "feb": 2, "mär": 3, "mae": 3, "mrz": 3, "apr": 4, "mai": 5,
    "jun": 6, "jul": 7, "aug": 8, "sep": 9, "okt": 10, "nov": 11, "dez": 12,
}
DATE = r"(\d{1,2})\.(\d{1,2})\.(\d{4}|\d{2})\b"
DATE_RANGE_RE = re.compile(DATE + r"\s*-\s*" + DATE)
MONTH_RANGE_RE = re.compile(
    r"\b([A-Za-zä]{3,5})\.?\s*(\d{4}|\d{2})\s*-\s*([A-Za-zä]{3,5})\.?\s*(\d{4}|\d{2})\b"
)
STORNO_RE = re.compile(r"Storno zum\s*" + DATE)
PERCENTAGE_RE = re.compile(r"(\d+(?:[.,]\d+)?)\s*%")
WEEKLY_HOURS_RE = re.compile(r"(\d+(?:[.,]\d+)?)\s*SWS")
PAYROLL_RE = re.compile(r"^VIVA (\d{2})(\d{4})\b")
# Payroll texts end with the date of birth, which we do not want to keep.
PAYROLL_BIRTHDATE_RE = re.compile(r"^(VIVA \d{6} .*?)\s+\d{6,8}\s*$")
PERSON_HEADER_RE = re.compile(r"^\s*([A-Za-z/]+)\s*[,;]\s*(.+?)\s*,\s*(.+?)\s*$")
CONTRACT_TYPE_RE = re.compile(r"^\s*(AV|SHK|WHK|TUT)\b")


@dataclass
class ParsedPosition:
    reference: str
    kind: str
    cost_type: str
    title: str = ""
    description: str = ""
    person_name: str = ""
    contract_type: str = ""
    contract_periods: list = field(default_factory=list)
    percentage: Decimal | None = None
    weekly_hours: Decimal | None = None
    first_date: date | None = None
    last_date: date | None = None
    actual: Decimal = Decimal("0.00")
    commitment: Decimal = Decimal("0.00")
    monthly_actuals: dict = field(default_factory=dict)
    bookings: list = field(default_factory=list)


@dataclass
class ParsedExport:
    psp_element: str
    row_count: int
    first_booking: date | None
    last_booking: date | None
    cost_types: dict
    positions: list


def is_gm_export(path):
    try:
        headers, _ = _read_rows(path, header_only=True)
    except Exception:
        return False
    return REQUIRED_HEADERS <= set(headers)


def parse_gm_export(path):
    headers, rows = _read_rows(path)
    missing = REQUIRED_HEADERS - set(headers)
    if missing:
        raise ValueError(f"Fehlende Spalten im SAP-Export: {', '.join(sorted(missing))}")

    details = [row for row in rows if _text(row.get("Fonds"))]
    if not details:
        raise ValueError("Der SAP-Export enthält keine Buchungszeilen.")

    psp_elements = {_text(row["PSP-Element"]) for row in details} - {""}
    if len(psp_elements) != 1:
        raise ValueError(
            "Der SAP-Export muss genau ein PSP-Element enthalten, gefunden: "
            + (", ".join(sorted(psp_elements)) or "keines")
        )

    unknown = {
        _text(row["FMM-Werttyp"]) for row in details
    } - BUDGET_VALUE_TYPES - COMMITMENT_VALUE_TYPES - ACTUAL_VALUE_TYPES
    if unknown:
        raise ValueError(f"Unbekannte FMM-Werttypen im SAP-Export: {', '.join(sorted(unknown))}")

    cost_types = defaultdict(lambda: {"budget": Decimal("0"), "actual": Decimal("0"), "commitment": Decimal("0")})
    groups = defaultdict(list)
    parents = {}
    for row in details:
        reference = _text(row["Referenzbelegnummer"])
        predecessor = _strip_zeros(row["Vorg. Referenzschl."])
        if predecessor and predecessor != reference:
            parents.setdefault(reference, predecessor)

    for row in details:
        value_type = _text(row["FMM-Werttyp"])
        amount = _decimal(row["Transaktionswährung"])
        cost_type = _text(row["E/A-Art"])
        if value_type in BUDGET_VALUE_TYPES:
            cost_types[cost_type]["budget"] += amount
            continue
        bucket = "actual" if value_type in ACTUAL_VALUE_TYPES else "commitment"
        cost_types[cost_type][bucket] += amount
        start = _strip_zeros(row["Vorg. Referenzschl."]) or _text(row["Referenzbelegnummer"])
        groups[_root(start, parents)].append(row)

    booking_dates = [d for d in (_date(row["Buchungsdatum"]) for row in details) if d]
    return ParsedExport(
        psp_element=psp_elements.pop(),
        row_count=len(details),
        first_booking=min(booking_dates, default=None),
        last_booking=max(booking_dates, default=None),
        cost_types={
            key: {name: _decimal_string(value) for name, value in values.items()}
            for key, values in sorted(cost_types.items())
        },
        positions=[_build_position(reference, group) for reference, group in sorted(groups.items())],
    )


@transaction.atomic
def import_gm_export(path, file_name=None, user=None):
    parsed = parse_gm_export(path)
    try:
        fund = SAPFund.objects.get(fund_number=parsed.psp_element)
    except SAPFund.DoesNotExist as error:
        raise ValueError(
            f"Für das PSP-Element {parsed.psp_element} ist kein SAP-Fonds angelegt. "
            "Bitte im Admin beim Projekt als Fondsnummer eintragen."
        ) from error

    SAPImport.objects.filter(fund=fund).delete()
    sap_import = SAPImport.objects.create(
        fund=fund,
        file_name=file_name or Path(path).name,
        imported_by=user,
        row_count=parsed.row_count,
        first_booking=parsed.first_booking,
        last_booking=parsed.last_booking,
        cost_types=parsed.cost_types,
    )
    SAPPosition.objects.bulk_create(
        SAPPosition(sap_import=sap_import, **position.__dict__) for position in parsed.positions
    )
    return sap_import


def parse_contract_texts(texts):
    """Extract contract periods, percentage and weekly hours from reservation texts."""
    periods = []
    storno_dates = []
    percentage = None
    weekly_hours = None
    for text in texts:
        for match in DATE_RANGE_RE.finditer(text):
            start = _parse_date(*match.groups()[:3])
            end = _parse_date(*match.groups()[3:])
            if start and end and start <= end:
                periods.append([start, end])
        if not DATE_RANGE_RE.search(text):
            for match in MONTH_RANGE_RE.finditer(text):
                start_month = _month(match[1])
                end_month = _month(match[3])
                if start_month and end_month:
                    start = date(_year(match[2]), start_month, 1)
                    end = _month_end(_year(match[4]), end_month)
                    if start <= end:
                        periods.append([start, end])
        for match in STORNO_RE.finditer(text):
            storno = _parse_date(*match.groups())
            if storno:
                storno_dates.append(storno)
        if percentage is None and (match := PERCENTAGE_RE.search(text)):
            percentage = _decimal(match[1].replace(",", "."))
        if weekly_hours is None and (match := WEEKLY_HOURS_RE.search(text)):
            weekly_hours = _decimal(match[1].replace(",", "."))

    if storno_dates:
        storno = min(storno_dates)
        periods = [[start, min(end, storno) if start <= storno else end] for start, end in periods]

    merged = []
    for start, end in sorted(periods):
        if merged and start <= merged[-1][1]:
            merged[-1][1] = max(merged[-1][1], end)
        else:
            merged.append([start, end])
    return merged, percentage, weekly_hours


def _build_position(reference, rows):
    rows = sorted(rows, key=lambda row: (_date(row["Buchungsdatum"]) or date.min, _text(row["Referenzbelegnummer"])))
    origin_rows = [row for row in rows if _text(row["Referenzbelegnummer"]) == reference] or rows
    cost_type_counts = Counter(_text(row["E/A-Art"]) for row in origin_rows)
    cost_type = cost_type_counts.most_common(1)[0][0]
    title = next((_text(row["Belegkopftext"]) for row in origin_rows if _text(row["Belegkopftext"])), "")
    texts = [_clean_text(row["Text"]) for row in origin_rows if _text(row["Text"])]
    description = texts[-1] if texts else ""
    if not title:
        title = description or next((_clean_text(row["Text"]) for row in rows if _text(row["Text"])), "")

    kind = _classify(rows, cost_type, [title, description])

    position = ParsedPosition(reference=reference, kind=kind, cost_type=cost_type, title=title[:255], description=description)
    dates = [d for d in (_date(row["Buchungsdatum"]) for row in rows) if d]
    position.first_date = min(dates, default=None)
    position.last_date = max(dates, default=None)

    monthly = defaultdict(Decimal)
    for row in rows:
        value_type = _text(row["FMM-Werttyp"])
        amount = _decimal(row["Transaktionswährung"])
        if value_type in ACTUAL_VALUE_TYPES:
            position.actual += amount
            if kind == SAPPositionKind.STAFF:
                monthly[_salary_month(row)] += amount
        else:
            position.commitment += amount
        position.bookings.append({
            "date": _date_string(row["Buchungsdatum"]),
            "value_type": value_type,
            "commitment": value_type in COMMITMENT_VALUE_TYPES,
            "cost_type": _text(row["E/A-Art"]),
            "document": _text(row["Referenzbelegnummer"]),
            "partner": _text(row.get("Name 1")),
            "text": _clean_text(row["Text"]),
            "amount": _decimal_string(amount),
        })

    if kind == SAPPositionKind.STAFF:
        position.monthly_actuals = {
            month: _decimal_string(amount) for month, amount in sorted(monthly.items()) if amount
        }
        header = PERSON_HEADER_RE.match(title)
        if header:
            position.person_name = f"{header[2]}, {header[3]}"
        type_match = CONTRACT_TYPE_RE.match(description) or (header and re.match(r"(AV|SHK|WHK|TUT)", header[1]))
        position.contract_type = type_match[1] if type_match else ""
        periods, position.percentage, position.weekly_hours = parse_contract_texts(
            _clean_text(row["Text"]) for row in origin_rows if _text(row["FMM-Werttyp"]) in COMMITMENT_VALUE_TYPES
        )
        position.contract_periods = [[start.isoformat(), end.isoformat()] for start, end in periods]

    position.actual = position.actual.quantize(Decimal("0.01"))
    position.commitment = position.commitment.quantize(Decimal("0.01"))
    return position


def _classify(rows, cost_type, texts):
    budget_groups = {_budget_group(row) for row in rows} - {""}
    all_cost_types = {_text(row["E/A-Art"]) for row in rows}
    # Travel (RKE/RKA) is sometimes booked on personnel cost types; the text decides.
    if any(TRAVEL_TEXT_RE.match(text) for text in texts):
        return SAPPositionKind.TRAVEL
    if STAFF_BUDGET_GROUP in budget_groups or all_cost_types & STAFF_COST_TYPES:
        return SAPPositionKind.STAFF
    if budget_groups == {INCOME_BUDGET_GROUP} or all_cost_types <= INCOME_COST_TYPES:
        return SAPPositionKind.INCOME
    if cost_type in TRAVEL_COST_TYPES:
        return SAPPositionKind.TRAVEL
    return SAPPositionKind.OTHER


def _budget_group(row):
    """First digit of the budget title in the Finanzposition, e.g. "4" for 1539.41.42941."""
    return _text(row.get("Finanzposition")).rsplit(".", 1)[-1][:1]


def _root(reference, parents):
    seen = set()
    while reference in parents and reference not in seen:
        seen.add(reference)
        reference = parents[reference]
    return reference


def _salary_month(row):
    match = PAYROLL_RE.match(_text(row["Text"]))
    if match:
        return f"{match[2]}-{match[1]}"
    period = min(max(int(_text(row["Buchungsperiode"]) or 1), 1), 12)
    return f"{_text(row['Geschäftsjahr'])}-{period:02d}"


def _read_rows(path, header_only=False):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = workbook.worksheets[0].iter_rows(values_only=True)
        try:
            headers = [_text(value) for value in next(rows)]
        except StopIteration as error:
            raise ValueError("Leerer SAP-Export.") from error
        if header_only:
            return headers, []
        return headers, [dict(zip(headers, values)) for values in rows]
    finally:
        workbook.close()


def _clean_text(value):
    text = _text(value)
    match = PAYROLL_BIRTHDATE_RE.match(text)
    return match[1] if match else text


def _parse_date(day, month, year):
    try:
        return date(_year(year), int(month), int(day))
    except ValueError:
        return None


def _year(value):
    year = int(value)
    return year + 2000 if year < 100 else year


def _month(name):
    return MONTHS.get(name.lower()[:3])


def _month_end(year, month):
    return date(year, month, monthrange(year, month)[1])


def _strip_zeros(value):
    return _text(value).lstrip("0")


def _text(value):
    return "" if value is None else str(value).strip()


def _decimal(value):
    if value in (None, ""):
        return Decimal("0")
    return Decimal(str(value))


def _decimal_string(value):
    return format(Decimal(value).quantize(Decimal("0.01")), "f")


def _date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def _date_string(value):
    parsed = _date(value)
    return parsed.isoformat() if parsed else None
